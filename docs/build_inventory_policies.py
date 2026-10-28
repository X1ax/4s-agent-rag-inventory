"""Build the controlled inventory policy renditions using the document runtime."""
import hashlib
import json
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data/knowledge/inventory"
CLAUSE = re.compile(r"INV-[A-Z]{3}-\d{3}-C\d{3}")


def read_source(path):
    _, header, body = path.read_text(encoding="utf-8").split("---", 2)
    metadata = {}
    for line in header.strip().splitlines():
        key, value = line.split(":", 1)
        metadata[key] = None if value.strip() == "null" else value.strip()
    return metadata, body.strip()


def blocks(body):
    lines = body.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("# "):
            index += 1
            continue
        if line.startswith("## "):
            yield "heading", line[3:]
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            yield "table", rows
            continue
        else:
            yield "paragraph", line
        index += 1


def control_rows(meta):
    return [("文档编号", meta["document_id"]), ("版本", meta["version"]),
        ("发布状态", "待企业核准  尚未生效"), ("起草日期", meta["drafted_on"]),
        ("归口岗位", meta["owner_role"]), ("批准与生效", "批准人待登记  生效日期待核准")]


def format_table(table, header=True):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn(f"w:{key}"), value)
                borders.append(edge)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:w"), "90")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            properties.append(margins)
            if header and row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EDEFF2")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.05
                snap = OxmlElement("w:snapToGrid")
                snap.set(qn("w:val"), "0")
                paragraph._p.get_or_add_pPr().append(snap)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    if header and row_index == 0:
                        run.bold = True
    if header:
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)


def write_docx(meta, body, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    section.header_distance = section.footer_distance = Inches(.3)
    for grid in section._sectPr.findall(qn("w:docGrid")):
        section._sectPr.remove(grid)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = f"{meta['document_id']}    库存制度受控资料    待核准"
    header.runs[0].font.size = Pt(9)
    doc.add_paragraph(meta["title"], style="Title")
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Inches(1.15), Inches(5.75)
    for key, value in control_rows(meta):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    format_table(table, header=False)
    doc.add_paragraph()
    for kind, content in blocks(body):
        if kind == "heading":
            doc.add_paragraph(content, style="Heading 1")
        elif kind == "paragraph":
            paragraph = doc.add_paragraph()
            match = CLAUSE.match(content)
            if match:
                paragraph.add_run(match[0] + " ").bold = True
                paragraph.add_run(content[len(match[0]):].lstrip())
            else:
                paragraph.add_run(content)
        else:
            table = doc.add_table(rows=0, cols=len(content[0]))
            table.autofit = False
            widths = [1.2] + [(6.9 - 1.2) / (len(content[0]) - 1)] * (len(content[0]) - 1)
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            for values in content:
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    cell.text = value
            format_table(table)
            doc.add_paragraph()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("待核准稿  第 ")
    for field in ("PAGE", "NUMPAGES"):
        if field == "NUMPAGES":
            footer.add_run(" 页 共 ")
        node = OxmlElement("w:fldSimple")
        node.set(qn("w:instr"), field)
        footer._p.append(node)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(9)
    doc.core_properties.title = meta["title"]
    doc.core_properties.author = "项目建设组"
    doc.core_properties.subject = f"{meta['document_id']} {meta['version']} 待企业核准"
    doc.save(output)


def write_pdf(meta, body, output):
    base = ParagraphStyle("Body", fontName="PolicyCJK", fontSize=10.5, leading=17,
        spaceAfter=8, wordWrap="CJK", alignment=TA_LEFT)
    heading = ParagraphStyle("Heading", parent=base, fontName="PolicyCJKBold", fontSize=13,
        leading=19, spaceBefore=11, spaceAfter=8, keepWithNext=True)
    title = ParagraphStyle("Title", parent=heading, fontSize=20, leading=28, spaceBefore=0)
    small = ParagraphStyle("Small", parent=base, fontSize=9, leading=14, spaceAfter=2)
    flow = [Paragraph(escape(meta["title"]), title), Spacer(1, 8)]
    metadata = Table([[Paragraph(escape(key), small), Paragraph(escape(value), small)]
        for key, value in control_rows(meta)], colWidths=[78, 418])
    metadata.setStyle(TableStyle([("GRID", (0,0), (-1,-1), .5, colors.HexColor("#d9d9d9")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("TOPPADDING", (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5)]))
    flow.extend([metadata, Spacer(1, 14)])
    for kind, content in blocks(body):
        if kind == "heading":
            flow.append(Paragraph(escape(content), heading))
        elif kind == "paragraph":
            flow.append(Paragraph(escape(content), base))
        else:
            widths = [80] + [(496 - 80) / (len(content[0]) - 1)] * (len(content[0]) - 1)
            table = Table([[Paragraph(escape(value), small) for value in row] for row in content],
                colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.5,colors.HexColor("#d9d9d9")),
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#edeff2")),
                ("VALIGN",(0,0),(-1,-1),"MIDDLE"), ("TOPPADDING",(0,0),(-1,-1),7),
                ("BOTTOMPADDING",(0,0),(-1,-1),7)]))
            flow.extend([table, Spacer(1, 9)])
    def page(canvas, document):
        canvas.saveState()
        canvas.setFont("PolicyCJK", 8)
        canvas.drawString(58, 766, f"{meta['document_id']}    库存制度受控资料    待核准")
        canvas.drawCentredString(306, 27, f"待核准稿    第 {document.page} 页")
        canvas.restoreState()
    document = SimpleDocTemplate(str(output), pagesize=letter, leftMargin=58, rightMargin=58,
        topMargin=46, bottomMargin=42, title=meta["title"], author="项目建设组")
    document.build(flow, onFirstPage=page, onLaterPages=page)


def main():
    pdfmetrics.registerFont(TTFont("PolicyCJK", "C:/Windows/Fonts/msyh.ttc", subfontIndex=0))
    pdfmetrics.registerFont(TTFont("PolicyCJKBold", "C:/Windows/Fonts/msyhbd.ttc", subfontIndex=0))
    documents = []
    sources = sorted((BASE / "policies").glob("INV-*")) + sorted((BASE / "authoring").glob("INV-*.md"))
    for source in sources:
        if source.suffix not in {".md", ".txt"}:
            continue
        meta, body = read_source(source)
        native = BASE / "policies" / f"{meta['document_id']}.{meta['native_format']}"
        if meta["native_format"] == "docx":
            write_docx(meta, body, native)
        elif meta["native_format"] == "pdf":
            write_pdf(meta, body, native)
        clauses = CLAUSE.findall(body)
        assert len(clauses) == len(set(clauses))
        documents.append({**meta, "source_type": "project_authored_policy_draft",
            "path": native.relative_to(BASE).as_posix(), "authoring_path": source.relative_to(BASE).as_posix(),
            "sha256": hashlib.sha256(native.read_bytes()).hexdigest(), "clause_ids": clauses,
            "citation_locator": "pdf_page_and_clause" if meta["native_format"] == "pdf" else "section_and_clause",
            "index_status": "not_indexed", "default_retrieval_enabled": False})
        if meta["native_format"] == "pdf":
            documents[-1]["clause_pages"] = {clause: index for index, page in enumerate(PdfReader(native).pages, 1)
                for clause in clauses if clause in re.sub(r"\s+", "", page.extract_text())}
    manifest = {"schema_version": 1, "collection": "inventory_policies", "language": "zh-CN",
        "drafted_on": "2026-10-22", "approval_status": "pending_approval",
        "release_note": "企业制度编制稿，未经实际经营主体核准；默认不得作为现行有效制度检索。",
        "canonical_source_rule": "Only documents[].path is an ingest source; authoring and QA files are excluded.",
        "documents": sorted(documents, key=lambda item: item["document_id"])}
    assert len(manifest["documents"]) == 12
    (BASE / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"documents": len(documents), "clauses": sum(len(row["clause_ids"]) for row in documents),
        "native_formats": {kind: sum(row["native_format"] == kind for row in documents)
            for kind in ("md", "txt", "pdf", "docx")}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
