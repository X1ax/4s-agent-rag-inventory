import re
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from pypdf import PdfReader, PdfWriter

from backend.app import knowledge as k
from test_knowledge import project


@pytest.fixture(scope="module")
def corpus():
    return k.parse_corpora()


def source(corpus, doc_id):
    doc = next(d for d in corpus["documents"] if d["document_id"] == doc_id)
    return deepcopy(doc), (k.ROOT / doc["path"]).read_bytes()


@pytest.mark.parametrize("doc_id", ["INV-PUR-003", "INV-TRF-006", "INV-SUP-012", "INV-RCV-004", "INV-WHS-007"])
def test_all_native_clause_text_and_tables_match_independent_authoring_oracle(corpus, doc_id):
    # The authoring file is a test oracle only; the runtime receives native bytes.
    doc, raw = source(corpus, doc_id)
    units, contexts = k.parse_document(doc, raw)
    original = (k.ROOT / f"data/knowledge/inventory/authoring/{doc_id}.md").read_bytes()
    _, lines = k.frontmatter(original)
    expected = {match[1]: match[2] for _, line in lines if (match := k.CLAUSE.fullmatch(line))}
    compact = lambda text: re.sub(r"\s+", "", text)
    assert [u["unit_id"] for u in units] == list(expected)
    for unit in units:
        assert compact(unit["unit_text"]) == compact(unit["unit_id"] + " " + expected[unit["unit_id"]])
        assert unit["parser_version"] == k.NATIVE_PARSER_VERSION
        assert "line_start" not in unit and "line_end" not in unit
    _, blocks = k.sections(lines, "md", doc["title"])
    expected_tables = [t for b in blocks for t in k.tables(b["lines"])]
    actual_tables = [t for c in contexts for t in c["tables"]]
    assert len(actual_tables) == len(expected_tables) == 1
    assert actual_tables[0]["headers"] == expected_tables[0]["headers"]
    assert actual_tables[0]["rows"] == expected_tables[0]["rows"]
    assert actual_tables[0]["ordinal"] == 2


def test_docx_locations_and_paragraph_table_order_are_preserved(corpus):
    for doc_id in ("INV-PUR-003", "INV-TRF-006", "INV-SUP-012"):
        doc, raw = source(corpus, doc_id)
        units, _ = k.parse_document(doc, raw)
        blocks = list(Document(BytesIO(raw)).iter_inner_content())
        for unit in units:
            assert unit["block_start"] == unit["block_end"]
            assert blocks[unit["block_start"]-1].text == unit["unit_text"]
            assert "page_start" not in unit
    doc, raw = source(corpus, "INV-SUP-012")
    _, contexts = k.parse_document(doc, raw)
    parent = next(c for c in contexts if "INV-SUP-012-C012" in c["unit_ids"])
    text = parent["context_text"]
    assert text.index("INV-SUP-012-C012") < text.index("| 参数组 |") < text.index("INV-SUP-012-C013")
    assert parent["tables"][0]["block_start"] == 21


def test_pdf_real_pages_cross_page_section_and_table_location(corpus):
    for doc_id in ("INV-RCV-004", "INV-WHS-007"):
        doc, raw = source(corpus, doc_id)
        units, contexts = k.parse_document(doc, raw)
        assert all(u["page_start"] == u["page_end"] == doc["clause_pages"][u["unit_id"]] for u in units)
        assert all("block_start" not in u for u in units)
        parent = next(c for c in contexts if f"{doc_id}-C008" in c["unit_ids"])
        assert parent["page_start"] == 1 and parent["page_end"] == 2
        table = next(t for c in contexts for t in c["tables"])
        assert table["page_start"] == table["page_end"] == 2
        assert len(table["rows"]) == 4
        assert not any("库存制度受控资料" in u["embedding_text"] or "待核准稿    第" in u["embedding_text"] for u in units)


def pdf_text_override(monkeypatch, texts):
    class Page:
        def __init__(self, text):
            self.text = text
        def extract_text(self):
            return self.text
    class Reader:
        is_encrypted = False
        def __init__(self, *args, **kwargs):
            self.pages = [Page(text) for text in texts]
    monkeypatch.setattr("pypdf.PdfReader", Reader)


def test_clause_continuation_across_pages_keeps_full_text_and_actual_range(corpus, monkeypatch):
    doc, raw = source(corpus, "INV-RCV-004")
    original_units, _ = k.parse_document(doc, raw)
    texts = [page.extract_text() for page in PdfReader(BytesIO(raw)).pages]
    first, second = [text.splitlines() for text in texts]
    continuation = first.pop()
    second.insert(2, continuation)
    pdf_text_override(monkeypatch, ["\n".join(first), "\n".join(second)])
    units, _ = k.parse_document(doc, raw)
    changed = next(u for u in units if u["unit_id"] == "INV-RCV-004-C008")
    original = next(u for u in original_units if u["unit_id"] == changed["unit_id"])
    assert changed["unit_text"] == original["unit_text"]
    assert changed["page_start"] == 1 and changed["page_end"] == 2


@pytest.mark.parametrize("edit,expected", [
    (lambda t: t.replace("登记结果\n入库单与流水可核查\n查询原请求后重试\n", ""), "native_table"),
    (lambda t: t.replace("放行条件", "未知表头"), "native_table"),
    (lambda t: t.replace("第 2 页", "第 9 页"), "pdf_page_profile"),
])
def test_pdf_unrecognized_table_or_page_profile_requires_review(corpus, monkeypatch, edit, expected):
    doc, raw = source(corpus, "INV-RCV-004")
    texts = [edit(p.extract_text()) for p in PdfReader(BytesIO(raw)).pages]
    pdf_text_override(monkeypatch, texts)
    with pytest.raises(k.KnowledgeError) as error:
        k.parse_document(doc, raw)
    assert error.value.state == "needs_review" and error.value.code == expected


@pytest.mark.parametrize("encrypted", [False, True])
def test_real_blank_or_encrypted_pdf_is_not_silently_indexed(corpus, encrypted):
    doc, _ = source(corpus, "INV-RCV-004")
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    if encrypted:
        writer.encrypt("test-only")
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(k.KnowledgeError) as error:
        k.parse_document(doc, stream.getvalue())
    assert error.value.state == "needs_review"
    assert error.value.code == ("encrypted_pdf" if encrypted else "pdf_no_text")


@pytest.mark.parametrize("kind", ["merged", "revision", "missing_table", "metadata"])
def test_docx_unsupported_or_inconsistent_content_requires_review(corpus, kind):
    doc, raw = source(corpus, "INV-SUP-012")
    original = Document(BytesIO(raw))
    if kind == "merged":
        original.tables[1].cell(1, 0).merge(original.tables[1].cell(1, 1))
    elif kind == "revision":
        original.element.body.append(OxmlElement("w:ins"))
    elif kind == "missing_table":
        element = original.tables[1]._tbl
        element.getparent().remove(element)
    else:
        original.tables[0].cell(1, 1).text = "unregistered-version"
    stream = BytesIO()
    original.save(stream)
    with pytest.raises(k.KnowledgeError) as error:
        k.parse_document(doc, stream.getvalue())
    assert error.value.state == "needs_review"


@pytest.mark.parametrize("doc_id", ["INV-SUP-012", "INV-RCV-004"])
def test_corrupt_native_bytes_are_not_replaced_with_authoring_files(corpus, doc_id):
    doc, _ = source(corpus, doc_id)
    with pytest.raises(k.KnowledgeError) as error:
        k.parse_document(doc, b"not-a-native-file-sensitive-sentinel")
    assert error.value.state == "needs_review" and "sentinel" not in error.value.detail


def test_pdf_page_registry_mismatch_is_rejected(corpus):
    doc, raw = source(corpus, "INV-RCV-004")
    doc["clause_pages"]["INV-RCV-004-C001"] = 2
    with pytest.raises(k.KnowledgeError) as error:
        k.parse_document(doc, raw)
    assert error.value.code == "pdf_page_mismatch"


def test_native_failure_is_reported_by_corpus_without_authoring_fallback_or_export(project):
    import json
    path = project / "data/knowledge/inventory/policies/INV-RCV-004.pdf"
    path.write_bytes(b"broken-native-pdf")
    manifest_path = project / "data/knowledge/inventory/manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    next(d for d in manifest["documents"] if d["document_id"] == "INV-RCV-004")["sha256"] = k.digest(path.read_bytes())
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    snapshot = k.parse_corpora(project)
    doc = next(d for d in snapshot["documents"] if d["document_id"] == "INV-RCV-004")
    assert doc["parse_status"] == "needs_review" and doc["error_code"] == "invalid_pdf"
    assert not any(u["document_id"] == doc["document_id"] for u in snapshot["units"])
    with pytest.raises(k.KnowledgeError):
        k.export_parsed(snapshot, project)


@pytest.mark.parametrize("pages", [{}, {"INV-RCV-004-C001": True}])
def test_incomplete_or_boolean_pdf_page_registry_requires_review(project, pages):
    import json
    path = project / "data/knowledge/inventory/manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    next(d for d in manifest["documents"] if d["document_id"] == "INV-RCV-004")["clause_pages"] = pages
    path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    doc = next(d for d in k.parse_corpora(project)["documents"] if d["document_id"] == "INV-RCV-004")
    assert doc["parse_status"] == "needs_review" and doc["error_code"] == "invalid_pdf_locator"


def test_runtime_never_reads_authoring_and_native_rule_version_changes_only_native_ids(project, monkeypatch):
    from pathlib import Path
    original = Path.read_bytes
    def guarded_read(path):
        assert "authoring" not in path.parts
        return original(path)
    monkeypatch.setattr(Path, "read_bytes", guarded_read)
    first = k.parse_corpora(project)
    assert first["summary"]["units"] == 242
    monkeypatch.setattr(k, "NATIVE_PARSER_VERSION", "controlled-office-v2")
    changed = k.parse_corpora(project)
    native_ids = {u["chunk_id"] for u in first["units"] if u["parser_version"] == "controlled-office-v1"}
    text_ids = {u["chunk_id"] for u in first["units"] if u["parser_version"] == k.PARSER_VERSION}
    new_ids = {u["chunk_id"] for u in changed["units"]}
    assert len(native_ids) == 87 and native_ids.isdisjoint(new_ids)
    assert text_ids <= new_ids and first["parse_id"] != changed["parse_id"]
