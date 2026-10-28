"""Local upload permissions and actual persisted retrieval with mock embeddings."""
import io
import shutil
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from backend.app.config import Settings
from backend.app.knowledge import CORPORA, ROOT, KnowledgeError, parse_corpora
from backend.app.knowledge_uploads import SubmissionStore, extract_blocks
from backend.app.main import create_app
from backend.app.retrieval import build_index
from backend.app.security import current_user
from test_retrieval import FakeEmbedding


@pytest.fixture(scope='module')
def uploads(tmp_path_factory):
    root = tmp_path_factory.mktemp('uploads')
    for relative, _ in CORPORA.values(): shutil.copytree(ROOT/relative, root/relative)
    settings = Settings(_env_file=None, mysql_user='', embedding_api_key='test-only')
    build_index(settings, root, embedder=FakeEmbedding(settings))
    app = create_app(settings, knowledge_root=root)
    user = SimpleNamespace(id=71, role='storekeeper')
    app.dependency_overrides[current_user] = lambda: user
    with TestClient(app) as client: yield app, client, root, user, settings


def submit(client, text='验收资料：报损须登记原因和来源单据。', **updates):
    return client.post('/api/knowledge/submissions', params={'title':'上传回归资料','knowledge_base':'inventory_policies','format':'md',**updates}, content=text.encode('utf-8'))


def test_pending_rejected_and_user_isolation_never_grant_retrieval(uploads):
    app, client, root, user, _ = uploads
    user.id, user.role = 71, 'storekeeper'
    response = submit(client)
    assert response.status_code == 201, response.text
    item = response.json()
    assert item['status'] == 'pending'
    assert not any(d['document_id'] == 'UPL-'+item['id'] for d in parse_corpora(root)['documents'])
    assert client.post(f"/api/knowledge/submissions/{item['id']}/review", json={'approved':True}).status_code == 403
    user.id = 72
    assert not any(r['id'] == item['id'] for r in client.get('/api/knowledge/submissions').json()['items'])
    user.id, user.role = 73, 'admin'
    response = client.post(f"/api/knowledge/submissions/{item['id']}/review", json={'approved':False,'note':'内容不完整，请补充适用条件'})
    assert response.json()['status'] == 'rejected'
    assert client.post(f"/api/knowledge/submissions/{item['id']}/review", json={'approved':True}).status_code == 409
    assert not (root/'data/knowledge/uploads/approved.json').exists()


def test_admin_approval_rebuilds_and_persists_real_local_retrieval(uploads, monkeypatch):
    app, client, root, user, settings = uploads
    user.role = 'storekeeper'
    item = submit(client).json()
    user.role = 'admin'
    monkeypatch.setattr('backend.app.knowledge_api.build_index', lambda s, r: build_index(s,r,embedder=FakeEmbedding(s)))
    response = client.post(f"/api/knowledge/submissions/{item['id']}/review", json={'approved':True,'note':'已核对原文和适用范围'})
    assert response.status_code == 200, response.text
    assert response.json()['status'] == 'indexed' and response.json()['reviewer_id'] == user.id
    runtime = app.state.knowledge
    result = runtime.read(lambda index: index.search('报损登记', 'inventory_policies', scope='draft_preview', document_ids=['UPL-'+item['id']], query_vector=[1.0]+[0.0]*1023))
    assert result['status'] == 'retrieved' and '报损须登记' in result['hits'][0]['unit_text']
    assert result['hits'][0]['line_start'] == 1
    released = runtime.read(lambda index: index.search('报损登记','inventory_policies',scope='released_only',document_ids=['UPL-'+item['id']],query_vector=[1.0]+[0.0]*1023))
    assert not released['hits']
    generation = runtime.index.report['generation_id']
    assert client.post(f"/api/knowledge/submissions/{item['id']}/review",json={'approved':True}).json()['status'] == 'indexed'
    assert runtime.index.report['generation_id'] == generation
    assert client.post(f"/api/knowledge/submissions/{item['id']}/review",json={'approved':False}).status_code == 409
    runtime.close(); runtime.start()
    assert 'UPL-'+item['id'] in {d['document_id'] for d in runtime.index.sources['documents']}


def test_index_failure_rolls_back_publication_and_allows_retry(uploads, monkeypatch):
    app, client, root, user, _ = uploads
    user.role = 'admin'
    item = submit(client, '索引失败回归资料，禁止提前用于检索。').json()
    previous = (root/'data/knowledge/uploads/approved.json').read_bytes()
    def fail(*args): raise KnowledgeError('mock_index_failure','模拟构建失败')
    monkeypatch.setattr('backend.app.knowledge_api.build_index',fail)
    response = client.post(f"/api/knowledge/submissions/{item['id']}/review",json={'approved':True})
    assert response.status_code == 503
    assert (root/'data/knowledge/uploads/approved.json').read_bytes() == previous
    assert next(r for r in client.get('/api/knowledge/submissions').json()['items'] if r['id'] == item['id'])['status'] == 'index_failed'
    assert not app.state.knowledge.maintenance and app.state.knowledge.index is not None
    assert 'UPL-'+item['id'] not in {d['document_id'] for d in parse_corpora(root)['documents']}
    monkeypatch.setattr('backend.app.knowledge_api.build_index', lambda s,r: build_index(s,r,embedder=FakeEmbedding(s)))
    assert client.post(f"/api/knowledge/submissions/{item['id']}/review",json={'approved':True}).json()['status'] == 'indexed'


def test_interrupted_publication_can_be_rejected_without_lingering_source(uploads):
    _, _, root, _, _ = uploads
    store = SubmissionStore(root)
    item = store.submit(71,'中断恢复资料','inventory_policies','txt','审核发布中断回归。'.encode('utf-8'))
    def interrupted(): raise SystemExit('simulated process interruption')
    with pytest.raises(SystemExit): store.review(item['id'],73,True,'',SimpleNamespace(rebuild=interrupted))
    assert next(r for r in store.items(71) if r['id']==item['id'])['status']=='approved'
    assert store.review(item['id'],73,False,'不予入库',None)['status']=='rejected'
    assert 'UPL-'+item['id'] not in {d['document_id'] for d in parse_corpora(root)['documents']}


def test_invalid_document_bytes_size_and_empty_titles_are_rejected(uploads):
    _, client, _, user, _ = uploads
    user.role = 'storekeeper'
    for updates, text in [({'format':'exe'},'text'), ({'format':'pdf'},'not a pdf'), ({'format':'docx'},'not a docx'), ({'title':'  '},'text'), ({},'')]:
        assert submit(client,text,**updates).status_code == 422
    assert submit(client,'a'*(10*1024*1024+1)).status_code == 413
    assert client.post('/api/knowledge/submissions',params={'title':'标题','knowledge_base':'inventory_policies','format':'txt'},content=b'\xff').status_code == 422


def test_pdf_and_docx_uploads_have_real_locators():
    from pypdf import PdfWriter
    from docx import Document
    doc = Document(); doc.add_paragraph('Receiving checks require source documents.')
    table = doc.add_table(rows=1,cols=2); table.cell(0,0).text='Item'; table.cell(0,1).text='Quantity'
    file = io.BytesIO(); doc.save(file)
    chunks = extract_blocks(file.getvalue(),'docx')
    assert chunks[0]['location']['block_start'] == 1 and 'Receiving checks' in chunks[0]['text']
    assert chunks[1]['location']['table_ordinal'] == 1 and 'Item | Quantity' in chunks[1]['text']
    # Actual PDF text and page numbers, independent of a renderer/font dependency.
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    pdf = PdfWriter(); page = pdf.add_blank_page(width=300,height=200)
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'):pdf._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 100 Td (Verify received quantities.) Tj ET')
    page[NameObject('/Contents')] = pdf._add_object(stream)
    raw = io.BytesIO(); pdf.write(raw)
    assert extract_blocks(raw.getvalue(),'pdf') == [{'text':'Verify received quantities.','location':{'page_start':1,'page_end':1}}]


@pytest.mark.parametrize('format,mime',[('PNG','image/png'),('JPEG','image/jpeg'),('WEBP','image/webp')])
def test_image_upload_decodes_reencodes_and_serves_opaque_keys(uploads,format,mime):
    _, client, _, user, _ = uploads
    raw = io.BytesIO(); Image.new('RGB',(24,16),'blue').save(raw,format=format)
    user.role = 'storekeeper'
    assert client.post('/api/media/images',content=raw.getvalue(),headers={'Content-Type':mime}).status_code == 403
    user.role = 'admin'
    response = client.post('/api/media/images',content=raw.getvalue(),headers={'Content-Type':mime})
    assert response.status_code == 201, response.text
    item = response.json(); assert len(item['image_key']) == 39
    preview = client.get(item['url']); assert preview.status_code == 200 and preview.headers['content-type'] == 'image/png'
    with Image.open(io.BytesIO(preview.content)) as image: assert image.size == (24,16)


def test_image_invalid_mime_magic_length_and_keys(uploads):
    _, client, _, user, _ = uploads; user.role = 'admin'
    assert client.post('/api/media/images',content=b'<svg/>',headers={'Content-Type':'image/svg+xml'}).status_code == 415
    assert client.post('/api/media/images',content=b'bad',headers={'Content-Type':'image/png'}).status_code == 422
    assert client.post('/api/media/images',content=b'x'*(5*1024*1024+1),headers={'Content-Type':'image/png'}).status_code == 413
    assert client.get('/api/media/images/not-a-key').status_code == 404

