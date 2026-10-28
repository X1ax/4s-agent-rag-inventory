"""Real HTTP/auth/model upload-chain verification in an isolated local root and test DB."""
import json
import multiprocessing
import os
import shutil
import socket
import subprocess
import sys
import time
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4


def fixtures(directory):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    directory.mkdir(parents=True, exist_ok=True)
    note = '本资料仅用于隔离环境链路验收，不是正式企业制度，不代表实际库存。'
    cases = [
        {'format':'md','title':'青栎周转箱交接核查说明','expected':['封签编号','交接时间','接收岗位','QL-479'],
         'question':'根据本地参考资料《青栎周转箱交接核查说明》，交接时登记哪三项信息，交接标记是什么？','locator':{'line_start':5}},
        {'format':'pdf','title':'霁蓝样件包异常暂存说明','expected':['苍岩待核架','JL-582'],
         'question':'根据本地参考资料《霁蓝样件包异常暂存说明》，外包装破损时的暂存点和异常登记代码分别是什么？','locator':{'page_start':2}},
        {'format':'docx','title':'紫杉循环箱接收核对说明','expected':['ZS-763','紫杉附件缺项记录'],
         'question':'根据本地参考资料《紫杉循环箱接收核对说明》，接收核对标识是什么，附件缺项时用什么记录名称？','locator':{'block_start':4}},
    ]
    (directory/'md.md').write_text('# '+cases[0]['title']+'\n\n'+note+'\n\n青栎周转箱交接必须登记封签编号、交接时间和接收岗位三项信息，交接标记为QL-479。\n',encoding='utf-8')
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    pdf=canvas.Canvas(str(directory/'pdf.pdf'),pagesize=(595,842))
    for page,lines in enumerate([[note,'本说明的具体异常暂存要求见第二页。'],
        ['霁蓝样件包发现外包装破损时，暂存点为苍岩待核架。','异常登记代码为JL-582，并通知资料管理员核对。','未核对前不得办理入库。']],1):
        pdf.setFont('STSong-Light',18);pdf.drawString(50,780,cases[1]['title'])
        pdf.setFont('STSong-Light',12)
        for i,line in enumerate(lines): pdf.drawString(50,730-i*28,line)
        pdf.drawString(50,45,f'第{page}页  合成验收资料');pdf.showPage()
    pdf.save()
    doc=Document();doc.styles['Normal'].font.name='Microsoft YaHei';doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.add_paragraph(cases[2]['title'],'Title');doc.add_paragraph(note);doc.add_paragraph('紫杉循环箱接收前按下表核对，不据此生成任何库存数量。')
    table=doc.add_table(rows=1,cols=2);table.style='Table Grid'
    table.rows[0].cells[0].text='核对项目';table.rows[0].cells[1].text='登记要求'
    for label,value in [('接收核对标识','ZS-763'),('附件缺项时的记录名称','紫杉附件缺项记录')]:
        cells=table.add_row().cells;cells[0].text=label;cells[1].text=value
    doc.save(directory/'docx.docx')
    (directory/'cases.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def serve(root, port, settings):
    import uvicorn
    from backend.app.main import create_app
    uvicorn.run(create_app(settings,knowledge_root=Path(root)),host='127.0.0.1',port=port,log_level='critical',access_log=False)


def main():
    import httpx
    from pydantic import SecretStr
    from sqlalchemy import delete, func, select
    from sqlalchemy.orm import Session
    from backend.app.config import ROOT, Settings
    from backend.app.db import make_engine
    from backend.app.knowledge import CORPORA, digest
    from backend.app.models import AgentProposal, AuditLog, Document, StockBalance, StockLedger, User
    from backend.app.retrieval import build_index
    from backend.app.security import passwords
    settings=Settings()
    assert settings.test_mysql_database=='fours_agent_test' and settings.test_mysql_user=='fours_test'
    settings=settings.model_copy(update={'mysql_database':settings.test_mysql_database,'mysql_user':settings.test_mysql_user,
        'mysql_password':settings.test_mysql_password,'jwt_secret':SecretStr(uuid4().hex+uuid4().hex)})
    stamp=datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')
    root=ROOT/'data/uploads'/('chain-'+stamp);root.mkdir()
    bundle=Path.home()/'.cache/codex-runtimes/codex-primary-runtime'
    node=bundle/'dependencies/node/bin/node.exe';python=bundle/'dependencies/python/python.exe'
    for kind in ('pdf','documents'):
        marker=bundle/f'plugins/openai-primary-runtime/plugins/{kind}/skills/{kind}/container_tools/mark_artifact_operation_started.mjs'
        subprocess.run([str(node),str(marker),'--operation-kind','create','--expected-output-count','1','--output-format','pdf' if kind=='pdf' else 'docx'],check=True,capture_output=True)
    artifact_env={k:v for k,v in os.environ.items() if k not in {'PYTHONHOME','PYTHONPATH','__PYVENV_LAUNCHER__','VIRTUAL_ENV'}}
    subprocess.run([str(python),'-I','-X','utf8',str(ROOT/'evals/verify_upload_chain.py'),'--fixtures',str(root/'inputs')],check=True,env=artifact_env)
    cases=json.loads((root/'inputs/cases.json').read_text(encoding='utf-8'))
    for relative,_ in CORPORA.values(): shutil.copytree(ROOT/relative,root/relative)
    report={'scope':'real_http_development_e2e_not_independent_accuracy_score','started_at':datetime.now(UTC).isoformat(),
        'root':str(root),'model':settings.model_id,'embedding_model':settings.embedding_model_id,'checks':[], 'cases':cases,'answers':[]}
    target=ROOT/'evals/runs'/('upload-chain-'+stamp+'.json')
    process=None;engine=make_engine(settings);user_ids=[]
    def record(name,passed,detail=None):
        report['checks'].append({'name':name,'passed':bool(passed),'detail':detail})
        target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        print(json.dumps({'check':name,'passed':bool(passed)},ensure_ascii=False),flush=True)
    def start_server():
        with socket.socket() as sock: sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
        child=multiprocessing.Process(target=serve,args=(str(root),port,settings),daemon=True);child.start()
        url=f'http://127.0.0.1:{port}'
        for _ in range(150):
            try:
                if httpx.get(url+'/api/health',timeout=.5).status_code==200:return child,url
            except httpx.HTTPError:pass
            if not child.is_alive():break
            time.sleep(.2)
        child.terminate();child.join(10);raise RuntimeError('隔离验证服务启动失败')
    def counts():
        with Session(engine) as session:return {model.__tablename__:session.scalar(select(func.count()).select_from(model)) for model in (Document,StockBalance,StockLedger,AgentProposal)}
    try:
        report['baseline_index']=build_index(settings,root)
        before=counts();credentials={};prefix='up'+uuid4().hex[:12]
        with Session(engine) as session,session.begin():
            for role in ('admin','storekeeper'):
                password=uuid4().hex;username=prefix+'_'+role
                user=User(username=username,display_name='上传链路验收'+role,role=role,password_hash=passwords.hash(password))
                session.add(user);session.flush();user_ids.append(user.id);credentials[role]={'username':username,'password':password}
        process,url=start_server();client=httpx.Client(base_url=url,timeout=310)
        tokens={}
        for role in credentials:
            response=client.post('/api/auth/login',json=credentials[role]);response.raise_for_status()
            tokens[role]={'Authorization':'Bearer '+response.json()['access_token']}
        record('actual_account_login',len(tokens)==2)
        def call(method,path,role='storekeeper',**kwargs):
            response=client.request(method,path,headers=tokens[role],**kwargs)
            if response.status_code>=400:raise RuntimeError(f'{method} {path}: HTTP{response.status_code} '+response.text[:300])
            return response.json()
        def chat(case,stage):
            thread=call('POST','/api/assistant/threads')
            response=call('POST',f"/api/assistant/threads/{thread['id']}/messages",json={'message':case['question'],
                'client_request_id':uuid4().hex,'knowledge_scope':'draft_preview','allow_web':False})
            answer=response['answer'];history=call('GET',f"/api/assistant/threads/{thread['id']}")
            report['answers'].append({'format':case['format'],'stage':stage,'question':case['question'],'thread_id':thread['id'],'title':history['title'],'answer':answer})
            return answer,thread['id'],history
        def verify_answer(case,answer,stage):
            text='\n'.join(c['text'] for c in answer.get('claims',[]))
            citations=[c for c in answer.get('citations',[]) if c['document_id']==case['document_id']]
            record(case['format']+'_'+stage+'_answer',all(term in text for term in case['expected']) and bool(citations),{'status':answer.get('answer_status'),'claims':text})
            located=False
            for citation in citations:
                fragment=call('GET','/api/knowledge/chunks/'+citation['chunk_id'],params={'generation_id':answer['generation_id'],
                    'knowledge_base':'inventory_policies','knowledge_scope':'draft_preview'})['fragment']
                if all(fragment.get(k)==v for k,v in case['locator'].items()) and all(term in fragment['excerpt'] for term in case['expected']):
                    case['citation']=citation;located=True
            record(case['format']+'_'+stage+'_original_locator',located,case['locator'])
        for case in cases:
            raw=(root/'inputs'/f"{case['format']}.{case['format']}").read_bytes()
            result=call('POST','/api/knowledge/submissions',params={'title':case['title'],'knowledge_base':'inventory_policies','format':case['format']},content=raw)
            case.update(submission_id=result['id'],document_id='UPL-'+result['id'],source_sha256=digest(raw),preview=result['preview'])
            record(case['format']+'_pending',result['status']=='pending')
            denied=client.post(f"/api/knowledge/submissions/{result['id']}/review",headers=tokens['storekeeper'],json={'approved':True})
            record(case['format']+'_nonadmin_review_denied',denied.status_code==403)
            items=call('GET','/api/knowledge/documents',params={'knowledge_base':'inventory_policies','knowledge_scope':'draft_preview'})['items']
            record(case['format']+'_pending_not_searchable',case['document_id'] not in {d['document_id'] for d in items})
        answer,_,_=chat(cases[0],'before_review')
        record('pending_assistant_does_not_know_canary',cases[0]['expected'][-1] not in json.dumps(answer,ensure_ascii=False)
            and not any(c['document_id'].startswith('UPL-') for c in answer.get('citations',[])))
        for case in cases:
            reviewed=call('POST',f"/api/knowledge/submissions/{case['submission_id']}/review",role='admin',json={'approved':True,'note':'隔离环境合成资料，已核对原文与定位'})
            record(case['format']+'_approved_indexed',reviewed['status']=='indexed')
            answer,tid,history=chat(case,'after_review');case['thread_id']=tid;case['generated_title']=history['title']
            verify_answer(case,answer,'after_review')
            record(case['format']+'_summary_title',history.get('title_mode')=='generated' and history['title'] != case['question'][:28] and len(history['title'])<=24,history['title'])
            released=call('POST','/api/knowledge/retrieve',json={'question':case['question'],'knowledge_base':'inventory_policies','knowledge_scope':'released_only','document_ids':[case['document_id']]})
            record(case['format']+'_not_formal_policy',not released['hits'])
        manual='上传资料引用核对'
        call('PATCH',f"/api/assistant/threads/{cases[0]['thread_id']}",json={'title':manual})
        generation=call('GET','/api/knowledge/status')['generation_id']
        for case in cases:
            citation=case.get('citation')
            if citation and citation['generation_id']!=generation:
                outdated=client.get('/api/knowledge/chunks/'+citation['chunk_id'],headers=tokens['storekeeper'],params={'generation_id':citation['generation_id'],'knowledge_base':'inventory_policies','knowledge_scope':'draft_preview'})
                record(case['format']+'_old_generation_rejected',outdated.status_code==409)
                answer,_,_=chat(case,'after_all_publications');verify_answer(case,answer,'after_all_publications')
        rejected=call('POST','/api/knowledge/submissions',params={'title':'驳回排除验收','knowledge_base':'inventory_policies','format':'txt'},content='合成驳回标记RX-928不得用于回答。'.encode('utf-8'))
        rejected=call('POST',f"/api/knowledge/submissions/{rejected['id']}/review",role='admin',json={'approved':False,'note':'合成资料不予入库'})
        items=call('GET','/api/knowledge/documents',params={'knowledge_base':'inventory_policies','knowledge_scope':'draft_preview'})['items']
        record('rejected_not_searchable',rejected['status']=='rejected' and 'UPL-'+rejected['id'] not in {d['document_id'] for d in items})
        client.close();process.terminate();process.join(10);process,url=start_server();client=httpx.Client(base_url=url,timeout=310)
        persisted=call('GET','/api/knowledge/submissions')['items']
        record('process_restart_keeps_review_status',all(next(r for r in persisted if r['id']==case['submission_id'])['status']=='indexed' for case in cases))
        for case in cases:
            history=call('GET',f"/api/assistant/threads/{case['thread_id']}")
            record(case['format']+'_restart_history_title',history['title']==(manual if case['format']=='md' else case['generated_title']) and bool(history['turns'][0]['answer']['citations']))
            citation=case.get('citation')
            if citation:
                fragment=call('GET','/api/knowledge/chunks/'+citation['chunk_id'],params={'generation_id':citation['generation_id'],'knowledge_base':'inventory_policies','knowledge_scope':'draft_preview'})['fragment']
                record(case['format']+'_restart_original_locator',all(fragment.get(k)==v for k,v in case['locator'].items()))
        answer,_,_=chat(cases[1],'after_process_restart');verify_answer(cases[1],answer,'after_process_restart')
        record('no_business_writes',counts()==before,counts())
        report['final_index']=call('GET','/api/knowledge/status')
        client.close()
    except Exception as error:
        report['failure']=str(error);record('run_completed',False,str(error))
    finally:
        if process is not None and process.is_alive():process.terminate();process.join(10)
        with Session(engine) as session,session.begin():
            session.execute(delete(AgentProposal).where(AgentProposal.user_id.in_(user_ids)))
            session.execute(delete(AuditLog).where(AuditLog.operator_id.in_(user_ids)))
            session.execute(delete(User).where(User.id.in_(user_ids)))
        engine.dispose()
        report['passed']=bool(report['checks']) and all(c['passed'] for c in report['checks'])
        report['finished_at']=datetime.now(UTC).isoformat()
        target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        print(json.dumps({'report':str(target),'passed':report['passed'],'checks':len(report['checks'])},ensure_ascii=False),flush=True)
    return 0 if report['passed'] else 1


if __name__=='__main__':
    if len(sys.argv)==3 and sys.argv[1]=='--fixtures':fixtures(Path(sys.argv[2]))
    else:raise SystemExit(main())
